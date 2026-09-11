from contextlib import contextmanager
from datetime import date, datetime, timezone
from io import BytesIO
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.api.context import context
from backend.api.monitoring import _traffic_monthly_workbook
from backend.auth import current_user
from backend.main import create_app
from backend.repository import Repository


def _camera(camera_id: str, directory_id: int | None):
    return SimpleNamespace(
        id=camera_id,
        directory_id=directory_id,
        modes_json='["people_flow"]',
    )


def test_monthly_report_groups_halls_and_preserves_blank_days(monkeypatch):
    cameras = [_camera("cam-a", 1), _camera("cam-b", 1), _camera("cam-c", None)]
    aggregates = [
        (1, date(2026, 9, 1), 5),
        (1, date(2026, 9, 3), 0),
        (None, date(2026, 9, 1), 2),
    ]

    class FakeResult:
        def all(self):
            return aggregates

    class FakeSession:
        scalars_index = 0

        def scalars(self, _statement):
            result = [cameras, [SimpleNamespace(id=1, name="滨湖")]][self.scalars_index]
            self.scalars_index += 1
            return result

        def execute(self, _statement):
            return FakeResult()

    @contextmanager
    def fake_scope():
        yield FakeSession()

    monkeypatch.setattr("backend.repository.session_scope", fake_scope)
    report = Repository().traffic_monthly(
        date(2026, 9, 1), datetime(2026, 9, 11, 4, tzinfo=timezone.utc)
    )

    assert len(report["days"]) == 30
    assert [row["hall_name"] for row in report["rows"]] == ["滨湖营业厅", "未分组营业厅"]
    assert report["rows"][0]["values"][:3] == [5, None, 0]
    assert report["rows"][0]["values"][11:] == [None] * 19
    assert report["rows"][0]["monthly_total"] == 5
    assert report["daily_totals"][:3] == [7, None, 0]
    assert report["daily_totals"][11:] == [None] * 19
    assert report["grand_total"] == 7


def test_monthly_report_returns_leap_month_without_flow_cameras(monkeypatch):
    class FakeSession:
        def scalars(self, _statement):
            return []

    @contextmanager
    def fake_scope():
        yield FakeSession()

    monkeypatch.setattr("backend.repository.session_scope", fake_scope)
    report = Repository().traffic_monthly(
        date(2024, 2, 1), datetime(2026, 9, 11, tzinfo=timezone.utc)
    )

    assert len(report["days"]) == 29
    assert report["rows"] == []
    assert report["daily_totals"] == [None] * 29
    assert report["grand_total"] == 0


def _sample_report():
    return {
        "month": "2026-09",
        "timezone": "Asia/Shanghai",
        "days": ["2026-09-01", "2026-09-02", "2026-09-03"],
        "rows": [{
            "directory_id": 1,
            "hall_name": "滨湖营业厅",
            "values": [3, None, 0],
            "monthly_total": 3,
        }],
        "daily_totals": [3, None, 0],
        "grand_total": 3,
    }


def test_monthly_workbook_keeps_empty_cells_and_totals():
    workbook = load_workbook(BytesIO(_traffic_monthly_workbook(_sample_report()).getvalue()))
    sheet = workbook["人流月报"]

    assert [cell.value for cell in sheet[1]] == ["营业厅", "1日", "2日", "3日", "月合计"]
    assert [cell.value for cell in sheet[2]] == ["滨湖营业厅", 3, None, 0, 3]
    assert [cell.value for cell in sheet[3]] == ["当月人流总计", 3, None, 0, 3]
    assert sheet.freeze_panes == "B2"


def test_monthly_api_validates_month_and_exports_excel(monkeypatch):
    class FakeRepository:
        def traffic_monthly(self, month):
            assert month == date(2026, 9, 1)
            return _sample_report()

    app = create_app()
    app.dependency_overrides[current_user] = lambda: SimpleNamespace(
        id=9, username="operator", display_name="值班员", role="user", enabled=True
    )
    original_repository = context.repository
    monkeypatch.setattr(context, "repository", FakeRepository())
    try:
        client = TestClient(app)
        invalid = client.get("/api/traffic/monthly?month=2026-9")
        future = client.get("/api/traffic/monthly?month=2999-01")
        valid = client.get("/api/traffic/monthly?month=2026-09")
        exported = client.post("/api/traffic/monthly/export", json={"month": "2026-09"})
    finally:
        context.repository = original_repository
        app.dependency_overrides.clear()

    assert invalid.status_code == 422
    assert future.status_code == 422
    assert valid.status_code == 200
    assert valid.json()["rows"][0]["values"] == [3, None, 0]
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    sheet = load_workbook(BytesIO(exported.content))["人流月报"]
    assert sheet.cell(3, 1).value == "当月人流总计"
