from __future__ import annotations

import asyncio
import os
import shutil
import time
from datetime import date as date_type, datetime, timezone
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query, WebSocket, WebSocketDisconnect
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.api.context import context
from backend.api.presenters import alert_public, analysis_public
from backend.schemas import (
    AlertBatchDelete,
    AlertExportRequest,
    TrafficMonthlyRequest,
    WebhookManualSend,
)
from backend.capabilities import CORE_CAPABILITIES
from backend.auth import admin_user, current_user, websocket_user


router = APIRouter()

_process_started_at = time.time()
_cpu_sample: tuple[int, int] | None = None
_task_sample: tuple[float, int] | None = None


def _host_snapshot() -> dict[str, Any]:
    """Read lightweight host/container telemetry without an extra dependency."""
    global _cpu_sample
    cpu_percent = 0.0
    memory_total = memory_used = 0
    try:
        values = [int(value) for value in Path("/proc/stat").read_text().splitlines()[0].split()[1:]]
        idle, total = values[3] + (values[4] if len(values) > 4 else 0), sum(values)
        if _cpu_sample:
            idle_delta, total_delta = idle - _cpu_sample[0], total - _cpu_sample[1]
            if total_delta > 0:
                cpu_percent = max(0.0, min(100.0, (1 - idle_delta / total_delta) * 100))
        _cpu_sample = (idle, total)
    except (OSError, ValueError, IndexError):
        try:
            cpu_percent = min(100.0, os.getloadavg()[0] / max(1, os.cpu_count() or 1) * 100)
        except (AttributeError, OSError):
            pass
    try:
        memory = {}
        for line in Path("/proc/meminfo").read_text().splitlines():
            key, value = line.split(":", 1)
            memory[key] = int(value.strip().split()[0]) * 1024
        memory_total = memory.get("MemTotal", 0)
        memory_used = max(0, memory_total - memory.get("MemAvailable", memory.get("MemFree", 0)))
    except (OSError, ValueError, IndexError):
        pass
    disk = shutil.disk_usage("/")
    return {
        "cpu_percent": round(cpu_percent, 1),
        "cpu_count": os.cpu_count() or 1,
        "memory_used": memory_used,
        "memory_total": memory_total,
        "memory_percent": round(memory_used / memory_total * 100, 1) if memory_total else 0,
        "disk_used": disk.used,
        "disk_total": disk.total,
        "disk_percent": round(disk.used / disk.total * 100, 1) if disk.total else 0,
    }


def _prometheus_snapshot(processed: int) -> dict[str, Any]:
    """Expose the same application signals used by the provisioned Grafana dashboard."""
    global _task_sample
    from prometheus_client import REGISTRY

    now = time.monotonic()
    task_rate = 0.0
    if _task_sample:
        elapsed = now - _task_sample[0]
        if elapsed > 0:
            task_rate = max(0.0, (processed - _task_sample[1]) / elapsed)
    _task_sample = (now, processed)

    vlm_calls = 0.0
    latency_sum = latency_count = 0.0
    for family in REGISTRY.collect():
        for sample in family.samples:
            if sample.name == "monitor_vlm_calls_total":
                vlm_calls += sample.value
            elif sample.name == "monitor_analysis_seconds_sum":
                latency_sum += sample.value
            elif sample.name == "monitor_analysis_seconds_count":
                latency_count += sample.value
    return {
        "task_rate": round(task_rate, 3),
        "vlm_calls_total": int(vlm_calls),
        "analysis_average_ms": round(latency_sum / latency_count * 1000, 1) if latency_count else 0,
    }


@router.get("/health")
async def health() -> dict[str, Any]:
    return {"status": "ok", "version": "1.0.0", **await context.require_runtime().status()}


@router.get("/api/dashboard", dependencies=[Depends(current_user)])
async def dashboard() -> dict[str, Any]:
    return {
        **context.repository.dashboard(),
        "runtime": await context.require_runtime().status(),
    }


@router.get("/api/alerts", dependencies=[Depends(current_user)])
async def alerts(
    limit: int = Query(default=100, ge=1, le=500),
    camera_id: str | None = None,
    mode: str | None = None,
    severity: str | None = None,
    date: date_type | None = None,
) -> list[dict[str, Any]]:
    rows = context.repository.list_alerts(limit, camera_id, mode, severity, date)
    return [alert_public(row, context.repository.webhook_deliveries(row.id)) for row in rows]


def _alert_workbook(rows: list[Any]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "告警记录"
    max_evidence = max((len(getattr(row, "evidences", []) or []) or bool(row.evidence_path) for row in rows), default=1)
    max_evidence = max(1, int(max_evidence))
    evidence_headers = [f"证据{index}" for index in range(1, max_evidence + 1)]
    sheet.append(["告警名称", "事件", "原因", *evidence_headers, "时间"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 48
    for index in range(max_evidence):
        sheet.column_dimensions[get_column_letter(4 + index)].width = 24
    time_column = 4 + max_evidence
    sheet.column_dimensions[get_column_letter(time_column)].width = 22
    for row_index, alert in enumerate(rows, start=2):
        created_at = context.repository._aware_utc(alert.created_at).astimezone(
            ZoneInfo("Asia/Shanghai")
        ).strftime("%Y-%m-%d %H:%M:%S")
        mode_name = CORE_CAPABILITIES.get(alert.mode, {}).get("name", alert.mode)
        alert_name = alert.alert_name or getattr(alert, "camera_name", alert.camera_id)
        sheet.append([alert_name, mode_name, alert.reason, *(["无证据"] * max_evidence), created_at])
        sheet.row_dimensions[row_index].height = 76
        evidence_rows = list(getattr(alert, "evidences", []) or [])
        evidence_paths = [item.evidence_path for item in evidence_rows]
        if not evidence_paths and alert.evidence_path:
            evidence_paths = [alert.evidence_path]
        for evidence_index, evidence_path in enumerate(evidence_paths):
            root = context.settings.evidence_dir.resolve()
            evidence = (root / Path(evidence_path).name).resolve()
            try:
                evidence.relative_to(root)
                if evidence.is_file():
                    image = ExcelImage(str(evidence))
                    image.width, image.height = 120, 90
                    column = 4 + evidence_index
                    sheet.add_image(image, f"{get_column_letter(column)}{row_index}")
                    sheet.cell(row_index, column).value = ""
            except (ValueError, OSError):
                pass
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@router.post("/api/alerts/export", dependencies=[Depends(current_user)])
async def export_alerts(payload: AlertExportRequest) -> StreamingResponse:
    alert_date = None
    if payload.date:
        try:
            alert_date = date_type.fromisoformat(payload.date)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail="日期格式必须为 YYYY-MM-DD") from exc
    rows = context.repository.list_alerts(
        None, mode=payload.mode, severity=payload.severity,
        alert_date=alert_date, alert_ids=payload.alert_ids,
    )
    filename = f"alerts-{payload.date or 'all'}.xlsx"
    workbook = await asyncio.to_thread(_alert_workbook, rows)
    return StreamingResponse(
        workbook,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/api/alert-webhook-targets", dependencies=[Depends(current_user)])
async def alert_webhook_targets() -> dict[str, Any]:
    return {"items": [
        {"id": row.id, "name": row.name, "enabled": row.enabled}
        for row in context.repository.list_webhook_targets(enabled_only=True)
    ]}


@router.post("/api/alerts/webhook-send", dependencies=[Depends(current_user)])
async def send_alerts_to_webhooks(payload: WebhookManualSend) -> dict[str, int]:
    try:
        return await context.require_runtime().alerts.manual_send(payload.alert_ids, payload.webhook_target_ids)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.delete("/api/alerts", dependencies=[Depends(admin_user)])
async def delete_alerts(
    before_days: int = Query(default=30, ge=1, le=365),
    severity: str | None = None,
) -> dict[str, Any]:
    runtime = context.require_runtime()
    result = await asyncio.to_thread(runtime.cleanup.run, before_days, severity)
    return {
        "deleted": result["deleted"],
        "evidence_removed": result["evidence_removed"],
        "cutoff": result["cutoff"],
    }


@router.get("/api/system-monitor", dependencies=[Depends(current_user)])
async def system_monitor() -> dict[str, Any]:
    runtime = await context.require_runtime().status()
    processed = int(runtime.get("workers", {}).get("processed", 0))
    dashboard_snapshot = context.repository.dashboard()
    return {
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "uptime_seconds": max(0, int(time.time() - _process_started_at)),
        "host": _host_snapshot(),
        "application": {
            **_prometheus_snapshot(processed),
            "processed": processed,
            "failures": int(runtime.get("workers", {}).get("failures", 0)),
            "online_cameras": int(dashboard_snapshot.get("online", 0)),
            "camera_count": int(dashboard_snapshot.get("cameras", 0)),
        },
        "runtime": runtime,
    }


@router.post("/api/alerts/batch-delete", dependencies=[Depends(admin_user)])
async def delete_selected_alerts(payload: AlertBatchDelete) -> dict[str, Any]:
    runtime = context.require_runtime()
    return await asyncio.to_thread(runtime.cleanup.delete_selected, payload.alert_ids)


@router.get("/api/analyses", dependencies=[Depends(admin_user)])
async def analyses(
    limit: int = Query(default=100, ge=1, le=500), camera_id: str | None = None
) -> list[dict[str, Any]]:
    return [analysis_public(row) for row in context.repository.list_analyses(limit, camera_id)]


@router.get("/api/traffic", dependencies=[Depends(current_user)])
async def traffic(
    camera_id: str | None = None, limit: int = Query(default=1440, ge=1, le=10000)
) -> list[dict[str, Any]]:
    return [
        {
            "camera_id": row.camera_id,
            "bucket_start": row.bucket_start,
            "current_count": row.current_count,
            "entered": row.entered,
            "exited": row.exited,
        }
        for row in context.repository.traffic(camera_id, limit)
    ]


@router.get("/api/traffic/summary", dependencies=[Depends(current_user)])
async def traffic_summary() -> dict[str, Any]:
    return context.repository.traffic_summary()


def _month_start(month: str) -> date_type:
    try:
        parsed = date_type.fromisoformat(f"{month}-01")
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="月份格式必须为 YYYY-MM") from exc
    if len(month) != 7 or parsed.strftime("%Y-%m") != month:
        raise HTTPException(status_code=422, detail="月份格式必须为 YYYY-MM")
    current = (
        datetime.now(timezone.utc)
        .astimezone(ZoneInfo("Asia/Shanghai"))
        .date()
        .replace(day=1)
    )
    if parsed > current:
        raise HTTPException(status_code=422, detail="不能查询未来月份")
    return parsed


@router.get("/api/traffic/monthly", dependencies=[Depends(current_user)])
async def traffic_monthly(
    month: str = Query(pattern=r"^\d{4}-(0[1-9]|1[0-2])$"),
) -> dict[str, Any]:
    return context.repository.traffic_monthly(_month_start(month))


def _traffic_monthly_workbook(report: dict[str, Any]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人流月报"
    headers = [
        "营业厅",
        *[f"{index}日" for index in range(1, len(report["days"]) + 1)],
        "月合计",
    ]
    sheet.append(headers)
    for row in report["rows"]:
        sheet.append([row["hall_name"], *row["values"], row["monthly_total"]])
    sheet.append(["当月人流总计", *report["daily_totals"], report["grand_total"]])

    header_fill = PatternFill("solid", fgColor="0B5FA5")
    total_fill = PatternFill("solid", fgColor="D9EAF7")
    border = Border(
        left=Side(style="thin", color="A8C4D8"),
        right=Side(style="thin", color="A8C4D8"),
        top=Side(style="thin", color="A8C4D8"),
        bottom=Side(style="thin", color="A8C4D8"),
    )
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    total_row = sheet.max_row
    total_column = sheet.max_column
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            if cell.row > 1:
                cell.alignment = Alignment(
                    horizontal="left" if cell.column == 1 else "center",
                    vertical="center",
                )
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)
        cell.fill = total_fill
    for cell in sheet.iter_cols(min_col=total_column, max_col=total_column):
        for item in cell:
            item.font = Font(bold=True, color="FFFFFF" if item.row == 1 else "000000")
            if item.row > 1:
                item.fill = total_fill
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 24
    for index in range(2, total_column):
        sheet.column_dimensions[get_column_letter(index)].width = 7
    sheet.column_dimensions[get_column_letter(total_column)].width = 12
    sheet.row_dimensions[1].height = 24
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@router.post("/api/traffic/monthly/export", dependencies=[Depends(current_user)])
async def export_traffic_monthly(payload: TrafficMonthlyRequest) -> StreamingResponse:
    month = _month_start(payload.month)
    report = context.repository.traffic_monthly(month)
    workbook = await asyncio.to_thread(_traffic_monthly_workbook, report)
    return StreamingResponse(
        workbook,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="traffic-{payload.month}.xlsx"'},
    )


@router.get("/api/runtime/workers", dependencies=[Depends(admin_user)])
async def workers() -> dict[str, Any]:
    return await context.require_runtime().status()


@router.websocket("/ws/events")
async def events_socket(websocket: WebSocket) -> None:
    if websocket_user(websocket) is None:
        await websocket.close(code=4401)
        return
    await websocket.accept()
    runtime = context.require_runtime()
    queue = runtime.event_bus.subscribe()
    try:
        while True:
            try:
                await websocket.send_json(await asyncio.wait_for(queue.get(), timeout=25))
            except asyncio.TimeoutError:
                await websocket.send_json({"type": "heartbeat"})
    except WebSocketDisconnect:
        pass
    finally:
        runtime.event_bus.unsubscribe(queue)
